"""Generate sample documents with deliberately unhelpful filenames.

Run from the backend folder:
    python scripts/make_samples.py

The filenames say nothing about the content, which is exactly the situation
FileMind is meant to solve. These files are reused for search evaluation later.
"""

import sys
from pathlib import Path

import docx
import pymupdf

SAMPLES_DIR = Path(__file__).resolve().parents[1] / "samples"

AGREEMENT_PAGES = [
    """SOFTWARE DEVELOPMENT AGREEMENT

This Software Development Agreement ("Agreement") is entered into on 3 March 2025 between Alpha Technologies Pvt. Ltd., a company incorporated in Bengaluru, India ("Client"), and Beta Solutions LLP, a software engineering firm based in Pune, India ("Developer").

1. SCOPE OF WORK
The Developer shall design, develop, test and deliver a cloud-based inventory management platform for the Client, including a web dashboard, a REST API and an Android mobile application. The detailed functional requirements are described in Schedule A.

2. PROJECT TIMELINE
Development will proceed in three milestones: (a) requirements and architecture, completed within four weeks; (b) core platform development, completed within twelve weeks; and (c) testing, deployment and handover, completed within four weeks. The Developer shall provide weekly progress reports to the Client's project manager.

3. ACCEPTANCE TESTING
Each milestone deliverable shall be subject to acceptance testing by the Client within ten business days of delivery. The Client shall notify the Developer in writing of any defects, which the Developer shall correct at no additional cost.""",
    """4. FEES AND PAYMENT
The total fee for the project is INR 18,00,000 (eighteen lakh rupees), payable in three instalments of 30%, 40% and 30% upon acceptance of each milestone. Invoices are payable within thirty days. Late payments shall attract interest at 1% per month.

5. INTELLECTUAL PROPERTY
Upon full payment, all intellectual property rights in the source code, designs and documentation created specifically for the Client shall transfer to the Client. The Developer retains ownership of its pre-existing libraries and tools, and grants the Client a perpetual, non-exclusive licence to use them as part of the delivered software.

6. CONFIDENTIALITY
Each party shall keep confidential all business, technical and financial information disclosed by the other party, and shall not disclose it to any third party without prior written consent. This obligation survives termination of this Agreement for three years.""",
    """7. WARRANTY
The Developer warrants that the delivered software will perform materially in accordance with Schedule A for ninety days after final acceptance. During this warranty period the Developer shall fix reported bugs free of charge.

8. TERMINATION
Either party may terminate this Agreement with thirty days' written notice. If the Client terminates for convenience, it shall pay for all work completed up to the termination date.

9. GOVERNING LAW
This Agreement is governed by the laws of India. Disputes shall be resolved by arbitration in Bengaluru.

Signed for Alpha Technologies Pvt. Ltd.: Priya Raman, Chief Technology Officer
Signed for Beta Solutions LLP: Arjun Mehta, Managing Partner""",
    """SCHEDULE A - FUNCTIONAL REQUIREMENTS

A1. Inventory tracking: the platform shall track stock levels for each product across all warehouses in real time, and record every stock movement with the user, time and reason.

A2. Purchase orders: users shall be able to create, approve and send purchase orders to suppliers. Orders above INR 5,00,000 require approval from a finance manager.

A3. Low-stock alerts: the system shall send email and in-app alerts when stock for a product falls below its configured reorder level.

A4. Barcode scanning: the Android application shall scan product barcodes to record goods received, stock transfers between warehouses and stock counts.

A5. Reporting: the web dashboard shall provide reports on stock valuation, slow-moving items, supplier lead times and monthly consumption, exportable to Excel and PDF.

A6. Access control: the platform shall support roles for administrators, warehouse staff and finance users, with permissions configurable by administrators.

A7. Integrations: the REST API shall allow the Client's existing accounting software to read purchase orders and stock valuation data.

A8. Performance: dashboard pages shall load within three seconds for up to 200 concurrent users.""",
]

CONTRACT_PAGES = [
    """EMPLOYMENT CONTRACT

This Employment Contract is made between Gamma Retail Stores Pvt. Ltd. ("Employer") and Ms. Kavya Nair ("Employee").

1. POSITION
The Employee is appointed as Senior Store Manager at the Employer's Kochi branch, reporting to the Regional Operations Head.

2. COMPENSATION
The Employee shall receive a gross annual salary of INR 9,60,000, paid monthly, together with a performance bonus of up to 10% of annual salary.

3. WORKING HOURS AND LEAVE
Normal working hours are 48 hours per week across six days. The Employee is entitled to 24 days of paid annual leave and 12 days of sick leave per year.

4. NOTICE PERIOD
Either party may end this contract with sixty days' written notice, or salary in lieu of notice.""",
]


def write_pdf(path: Path, pages: list[str]) -> None:
    doc = pymupdf.open()
    for text in pages:
        page = doc.new_page()
        # A negative return value means the text did not fit and was silently dropped.
        if page.insert_textbox(pymupdf.Rect(60, 60, 535, 790), text, fontsize=10.5) < 0:
            raise ValueError(f"Text overflowed a page in {path.name}")
    doc.save(path)


def write_maintenance_docx(path: Path) -> None:
    document = docx.Document()
    document.add_heading("Website Maintenance and Support Contract", level=0)
    document.add_paragraph(
        "This contract is between Delta Hospitality Group (\"Customer\") and Epsilon Web Services "
        "(\"Provider\") for the ongoing maintenance of the Customer's hotel booking website."
    )
    document.add_heading("Services", level=1)
    document.add_paragraph(
        "The Provider shall apply security patches, update the content management system and plugins, "
        "monitor uptime, perform weekly backups and fix broken pages or booking-form errors."
    )
    document.add_paragraph(
        "Critical incidents, such as the booking website being unavailable, will receive a response "
        "within two hours, 24 hours a day."
    )
    document.add_heading("Fees", level=1)
    document.add_paragraph("The Customer shall pay the following annual support fees, invoiced each April:")
    table = document.add_table(rows=4, cols=2)
    rows = [
        ("Service", "Annual fee"),
        ("Hosting and monitoring", "INR 1,20,000"),
        ("Security updates and backups", "INR 90,000"),
        ("Content changes (up to 40 hours)", "INR 60,000"),
    ]
    for r, (service, fee) in enumerate(rows):
        table.cell(r, 0).text = service
        table.cell(r, 1).text = fee
    document.add_heading("Term", level=1)
    document.add_paragraph(
        "This contract runs for twelve months and renews automatically unless either party gives "
        "forty-five days' notice before the renewal date."
    )
    document.save(path)


def write_proposal_docx(path: Path) -> None:
    document = docx.Document()
    document.add_heading("Proposal: AI Customer Support Assistant", level=0)
    document.add_heading("Background", level=1)
    document.add_paragraph(
        "Zeta Insurance receives around 12,000 customer emails each month about policy renewals, "
        "claim status and document requests. Response times currently average three days."
    )
    document.add_heading("Proposed solution", level=1)
    document.add_paragraph(
        "We propose building an AI assistant that classifies incoming emails, drafts replies using a "
        "large language model, and retrieves answers from the policy knowledge base using semantic search. "
        "Human agents review every draft before it is sent."
    )
    document.add_heading("Phases and budget", level=1)
    document.add_paragraph(
        "Phase 1 is a six-week proof of concept on historical emails. Phase 2 integrates the assistant "
        "with the support inbox. The estimated budget for both phases is INR 32,00,000."
    )
    document.save(path)


NOTES = """Project sync notes - Thursday
Attendees: Rohan, Meera, Sanjay

1. Office move: the team moves to the 4th floor on the 15th. Sanjay to book movers.
2. Vendor payment: approved an advance payment of ₹50,000 to Omega Printing for the brochure
   order. Remaining balance due on delivery. Meera to raise the purchase order.
3. Laptop refresh: six laptops are out of warranty; Rohan to collect quotes from two suppliers.
4. Training: first-aid training scheduled for next month, attendance is mandatory.

Action items
- Sanjay: movers booking by Monday
- Meera: purchase order for Omega Printing
- Rohan: laptop quotes by end of week
"""


def write_scanned_pdf(path: Path) -> None:
    doc = pymupdf.open()
    pixmap = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 60, 80), False)
    pixmap.clear_with(235)
    doc.new_page().insert_image(pymupdf.Rect(60, 60, 535, 760), stream=pixmap.tobytes("png"))
    doc.save(path)


def main() -> None:
    SAMPLES_DIR.mkdir(exist_ok=True)
    write_pdf(SAMPLES_DIR / "agreement_001.pdf", AGREEMENT_PAGES)
    write_pdf(SAMPLES_DIR / "contract_045.pdf", CONTRACT_PAGES)
    write_maintenance_docx(SAMPLES_DIR / "document_final.docx")
    write_proposal_docx(SAMPLES_DIR / "draft_v2.docx")
    (SAMPLES_DIR / "notes.txt").write_text(NOTES, encoding="utf-8")
    write_scanned_pdf(SAMPLES_DIR / "scan_0001.pdf")
    for path in sorted(SAMPLES_DIR.iterdir()):
        print(f"{path.name:22} {path.stat().st_size:>8,} bytes")


if __name__ == "__main__":
    sys.exit(main())
