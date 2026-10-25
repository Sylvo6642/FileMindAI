"""Builders that create small PDF/DOCX/TXT files in memory for tests."""

import io
import zipfile

import docx
import pymupdf


def make_pdf(pages: list[str]) -> bytes:
    """One PDF page per string; an empty string makes an empty page."""
    doc = pymupdf.open()
    for text in pages:
        page = doc.new_page()
        if text:
            page.insert_textbox(pymupdf.Rect(50, 50, 550, 800), text, fontsize=10)
    return doc.tobytes()


def make_image_only_pdf(page_count: int = 1) -> bytes:
    doc = pymupdf.open()
    pixmap = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 20, 20), False)
    pixmap.clear_with(180)
    png = pixmap.tobytes("png")
    for _ in range(page_count):
        doc.new_page().insert_image(pymupdf.Rect(50, 50, 400, 400), stream=png)
    return doc.tobytes()


def make_mixed_pdf() -> bytes:
    """Page 1 has text, page 2 is image-only."""
    doc = pymupdf.open(stream=make_pdf(["Text on the first page."]), filetype="pdf")
    doc.insert_pdf(pymupdf.open(stream=make_image_only_pdf(), filetype="pdf"))
    return doc.tobytes()


def make_encrypted_pdf() -> bytes:
    doc = pymupdf.open(stream=make_pdf(["Secret text."]), filetype="pdf")
    return doc.tobytes(encryption=pymupdf.PDF_ENCRYPT_AES_256, owner_pw="o", user_pw="u")


def make_docx(blocks: list) -> bytes:
    """blocks: ("h", text) heading, ("p", text) paragraph, ("t", [[cells]]) table."""
    document = docx.Document()
    for kind, content in blocks:
        if kind == "h":
            document.add_heading(content, level=1)
        elif kind == "p":
            document.add_paragraph(content)
        elif kind == "t":
            table = document.add_table(rows=len(content), cols=len(content[0]))
            for r, row in enumerate(content):
                for c, value in enumerate(row):
                    table.cell(r, c).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_corrupt_docx() -> bytes:
    """A zip that has the right file names but garbage XML inside."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("[Content_Types].xml", "not xml")
        archive.writestr("word/document.xml", "not xml either")
    return buffer.getvalue()
